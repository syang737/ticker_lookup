"""Append ticker and previous close to every row of an uploaded CSV.

The original columns pass through untouched; every row gets a status, so a
row that could not be priced says why instead of disappearing. The result is
written as .xlsx by default: opened in Excel, a CSV loses CUSIPs to scientific
notation (78471E109) and leading zeros, and only a typed cell prevents that.
"""

from __future__ import annotations

import csv
import io
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .cache import FigiCache
from .identifiers import CUSIP, classify
from .massive import Massive
from .openfigi import OpenFigi

ID_HEADERS = ("cusip", "isin", "ticker", "symbol", "identifier", "security id", "security")
OUT_COLUMNS = ["lookup_ticker", "close", "close_date", "lookup_status", "lookup_name", "lookup_figi"]


class InputError(ValueError):
    """The upload cannot be processed at all; the message is for the user."""


@dataclass
class Table:
    header: list[str]
    # Original cells are strings; close is a float and close_date a date.
    rows: list[list]
    id_col: int


@dataclass
class Summary:
    rows: int
    priced: int
    column: str
    close_date: date
    statuses: Counter = field(default_factory=Counter)


def decode(data: bytes) -> str:
    # utf-8-sig: Excel's "CSV UTF-8" starts with a byte-order mark.
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Plain "CSV" from Windows Excel.
        return data.decode("cp1252")


def detect_column(header: list[str], rows: list[list[str]]) -> int:
    lowered = [h.strip().lower() for h in header]
    for name in ID_HEADERS:
        for i, h in enumerate(lowered):
            if h == name or h.startswith(name):
                return i
    # No telling header: the column where most values look like identifiers.
    sample = rows[:200]
    scores = [sum(1 for r in sample if i < len(r) and classify(r[i])[1] not in (None, "TICKER")) for i in range(len(header))]
    if not scores or max(scores) == 0:
        raise InputError("No column of CUSIPs or ISINs found. Name the column 'CUSIP', 'ISIN' or 'Ticker'.")
    return scores.index(max(scores))


def massive_ticker(row: dict) -> str | None:
    # Massive prices equities only; a bond's FIGI "ticker" is a description.
    if row.get("marketSector") != "Equity" or not row.get("ticker"):
        return None
    # OpenFIGI writes share classes as BRK/B; Massive as BRK.B.
    return row["ticker"].replace("/", ".")


def choose(value: str, rows: list[dict], closes: dict[str, float]) -> tuple[dict | None, str | None, str]:
    """(chosen FIGI row, Massive ticker, status) for one identifier's matches."""
    if not rows:
        if value[3:6] in ("ESC", "CVR", "CNT"):
            return None, None, "not found (escrow / CVR placeholder)"
        return None, None, "not found"
    us = [r for r in rows if r.get("exchCode") == "US"]
    if not us:
        if len(rows) == 1 and not rows[0].get("exchCode"):
            r = rows[0]
            return r, None, f"not priceable ({r.get('marketSector')} {r.get('securityType')})"
        # Usually an old CUSIP after a corporate action: only foreign venues
        # remain, and their tickers (BZHGBP) would price something else.
        return rows[0], None, "no US listing"
    tradable = [(r, t) for r in us if (t := massive_ticker(r))]
    if not tradable:
        r = us[0]
        return r, None, f"not priceable ({r.get('marketSector')} {r.get('securityType')})"
    # A closed-end fund also maps to its NAV symbol, which never trades:
    # the candidate with a close is the listing.
    priced = [(r, t) for r, t in tradable if t in closes]
    if not priced:
        r, t = tradable[0]
        return r, t, "no close"
    r, t = priced[0]
    return r, t, "ok" if len(priced) == 1 else f"ok, ambiguous ({', '.join(t for _, t in priced)})"


def enrich(
    data: bytes,
    figi: OpenFigi,
    massive: Massive,
    cache: FigiCache | None = None,
) -> tuple[Table, Summary]:
    rows = list(csv.reader(io.StringIO(decode(data))))
    rows = [r for r in rows if any(c.strip() for c in r)]
    if len(rows) < 2:
        raise InputError("The file needs a header row and at least one data row.")
    header, body = rows[0], rows[1:]
    col = detect_column(header, body)

    classified = [classify(r[col] if col < len(r) else "") for r in body]
    keys = list(dict.fromkeys((v, k) for v, k, _ in classified if k))

    found = cache.get_many(keys) if cache else {}
    missing = [k for k in keys if k not in found]
    if missing:
        fresh = figi.map(missing)
        if cache:
            cache.put_many(fresh)
        found.update(fresh)

    close_date, closes = massive.previous_closes()

    table = Table(header=header + OUT_COLUMNS, rows=[], id_col=col)
    summary = Summary(rows=len(body), priced=0, column=header[col], close_date=close_date)
    for r, (value, kind, note) in zip(body, classified):
        if kind is None:
            match, ticker, status = None, None, f"invalid: {note}"
        else:
            match, ticker, status = choose(value if kind == CUSIP else "", found[(value, kind)], closes)
            if note:
                status += f", {note}"
        close = closes.get(ticker) if ticker else None
        summary.priced += close is not None
        summary.statuses[status.split(" (")[0].split(",")[0]] += 1
        table.rows.append(r + [
            ticker or "",
            close,
            close_date if close is not None else None,
            status,
            (match or {}).get("name") or "",
            (match or {}).get("compositeFIGI") or "",
        ])
    return table, summary


def to_csv(table: Table) -> bytes:
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(table.header)
    for r in table.rows:
        w.writerow(["" if v is None else v.isoformat() if isinstance(v, date) else str(v) for v in r])
    # utf-8-sig so Excel reads the names correctly, on Windows too.
    return out.getvalue().encode("utf-8-sig")


# Only values that survive the trip to a number and back unchanged become
# numbers: 0123 (an account), 1E5 or 12345678901234567 (past float precision)
# stay text.
NUMBER_RE = re.compile(r"^-?(0|[1-9]\d{0,14})(\.\d{1,15})?$")


def _cell(ws, value, text: bool = False) -> WriteOnlyCell:
    if isinstance(value, str) and not text and NUMBER_RE.match(value) and len(value) <= 16:
        value = float(value) if "." in value else int(value)
    cell = WriteOnlyCell(ws, value=value)
    if isinstance(value, str):
        # Without this, a value starting with "=" is written as a live formula.
        cell.data_type = "s"
        if text:
            cell.number_format = "@"
    elif isinstance(value, date):
        cell.number_format = "yyyy-mm-dd"
    return cell


def to_xlsx(table: Table) -> bytes:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Prices")
    ws.freeze_panes = "A2"
    for i, name in enumerate(table.header):
        width = max([len(name), *(len(str(r[i])) for r in table.rows[:200] if i < len(r))])
        ws.column_dimensions[get_column_letter(i + 1)].width = min(max(width + 2, 8), 45)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(table.header))}{len(table.rows) + 1}"

    bold = Font(bold=True)
    head = []
    for name in table.header:
        c = _cell(ws, name, text=True)
        c.font = bold
        head.append(c)
    ws.append(head)

    n_orig = len(table.header) - len(OUT_COLUMNS)
    # Text columns among those appended: everything but close and close_date.
    text_cols = {table.id_col} | {n_orig + i for i, c in enumerate(OUT_COLUMNS) if c not in ("close", "close_date")}
    for r in table.rows:
        ws.append([_cell(ws, v, text=i in text_cols) for i, v in enumerate(r)])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

